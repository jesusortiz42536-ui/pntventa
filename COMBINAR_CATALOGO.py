"""
COMBINAR CATALOGO CON ARCHIVOS DEL USUARIO
Lee los 2 archivos Excel del usuario, cruza contra la BD,
actualiza productos existentes y agrega nuevos.
Genera CATALOGO_COMPLETO_COMBINADO.xlsx
"""
import sqlite3
import os
import re
import random
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "farmacia.db")
ARCHIVO1 = r'C:\Users\chule\Downloads\LISTA FARMACIA 06MAYO2024 (1).xlsx'
ARCHIVO2 = r'C:\Users\chule\Downloads\listaz enero 2025.xlsx'
OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                      "CATALOGO_COMPLETO_COMBINADO.xlsx")

random.seed(777)

# Estilos
AZUL = "4A90E2"
VERDE = "27AE60"
NARANJA = "E67E22"
ROJO = "E74C3C"
BLANCO = "FFFFFF"
GRIS = "F2F2F2"
HEADER_FILL = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color=BLANCO, size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color="2C3E50", size=16)
SUBTITLE_FONT = Font(name="Calibri", bold=True, color="2C3E50", size=12)
NORMAL_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", bold=True, size=10)
VERDE_FILL = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
ROJO_FILL = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
GRIS_FILL = PatternFill(start_color=GRIS, end_color=GRIS, fill_type="solid")
AMARILLO_FILL = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
MONEY_FMT = '#,##0.00'
PCT_FMT = '0.0%'

MARGENES = {
    "precio_fahorro":     (1.05, 1.20),
    "precio_guadalajara": (1.10, 1.25),
    "precio_sanpablo":    (1.15, 1.35),
    "precio_benavides":   (1.12, 1.28),
    "precio_walmart":     (1.03, 1.15),
    "precio_sams":        (0.95, 1.10),
    "precio_moderna":     (1.02, 1.12),
}


def limpiar_nombre(nombre):
    """Normaliza un nombre de producto para comparacion."""
    if not nombre:
        return ""
    n = str(nombre).strip()
    n = n.replace('\n', ' ').replace('\r', ' ').replace('\xa0', ' ')
    n = re.sub(r'\s+', ' ', n).strip()
    return n.upper()


def extraer_clave_busqueda(nombre):
    """Genera clave para buscar coincidencias."""
    n = limpiar_nombre(nombre)
    # Quitar caracteres especiales
    n = re.sub(r'[^A-Z0-9\s/]', '', n)
    # Tomar las primeras 3 palabras significativas
    palabras = [p for p in n.split() if len(p) > 1 and p not in ('DE', 'CON', 'MG', 'ML', 'TAB', 'CX', 'GR')]
    return ' '.join(palabras[:3])


def extraer_sustancia_de_nombre(nombre):
    """Intenta extraer la sustancia activa del nombre comercial."""
    n = limpiar_nombre(nombre)
    # Buscar patron (SUSTANCIA)
    m = re.search(r'\(([A-Z/\s]+)\)', n)
    if m:
        return m.group(1).strip()
    return ""


def generar_precios_multi(precio_base):
    """Genera precios estimados para las 8 farmacias."""
    ps = float(precio_base)
    precios = {"precio_similares": ps}
    for col, (lo, hi) in MARGENES.items():
        precios[col] = round(ps * random.uniform(lo, hi), 2)
    return precios


def leer_archivo1():
    """Lee LISTA FARMACIA 06MAYO2024 - Hoja1
    Columnas: SUSTANCIA ACTIVA | NOMBRE COMERCIAL | EXISTENCIAS | PRECIO FARMACIA
    """
    print("  Leyendo LISTA FARMACIA 06MAYO2024...")
    wb = load_workbook(ARCHIVO1, data_only=True)
    ws = wb['Hoja1']
    productos = []
    for i in range(2, ws.max_row + 1):
        sustancia = ws.cell(row=i, column=1).value
        nombre = ws.cell(row=i, column=2).value
        stock = ws.cell(row=i, column=3).value
        precio = ws.cell(row=i, column=4).value
        if nombre and precio:
            nombre_limpio = limpiar_nombre(nombre)
            sustancia_limpia = limpiar_nombre(sustancia) if sustancia else ""
            try:
                precio_f = float(precio)
            except (ValueError, TypeError):
                continue
            try:
                stock_i = int(stock) if stock else 0
            except (ValueError, TypeError):
                stock_i = 0
            productos.append({
                "nombre": nombre_limpio,
                "sustancia": sustancia_limpia,
                "stock": stock_i,
                "precio": precio_f,
                "origen": "LISTA_MAYO2024",
            })
    wb.close()
    print(f"    {len(productos)} productos leidos")
    return productos


def leer_archivo2():
    """Lee listaz enero 2025 - multiples hojas."""
    print("  Leyendo listaz enero 2025...")
    wb = load_workbook(ARCHIVO2, data_only=True)
    productos = []
    vistos = set()

    # Hoja Factura: col B=articulo, C=descripcion, D=cant, E=precio unitario
    ws = wb['Factura']
    for i in range(9, ws.max_row + 1):
        nombre = ws.cell(row=i, column=3).value
        precio = ws.cell(row=i, column=5).value
        cant = ws.cell(row=i, column=4).value
        if nombre and precio:
            nombre_limpio = limpiar_nombre(nombre)
            if nombre_limpio in vistos or len(nombre_limpio) < 3:
                continue
            try:
                precio_f = float(precio)
            except (ValueError, TypeError):
                continue
            if precio_f <= 0:
                continue
            vistos.add(nombre_limpio)
            productos.append({
                "nombre": nombre_limpio,
                "sustancia": "",
                "stock": int(cant) if cant else 1,
                "precio": precio_f,
                "origen": "LISTA_ENERO2025_FACTURA",
            })

    # Hoja1: col A=codigo barras, B=nombre, C=cantidad, D=precio
    ws = wb['Hoja1']
    for i in range(1, ws.max_row + 1):
        nombre = ws.cell(row=i, column=2).value
        precio = ws.cell(row=i, column=4).value
        cant = ws.cell(row=i, column=3).value
        codigo = ws.cell(row=i, column=1).value
        if nombre and precio:
            nombre_limpio = limpiar_nombre(nombre)
            if nombre_limpio in vistos or len(nombre_limpio) < 3:
                continue
            try:
                precio_f = float(precio)
            except (ValueError, TypeError):
                continue
            if precio_f <= 0:
                continue
            vistos.add(nombre_limpio)
            productos.append({
                "nombre": nombre_limpio,
                "sustancia": "",
                "stock": int(cant) if cant else 1,
                "precio": precio_f,
                "origen": "LISTA_ENERO2025_HOJA1",
                "ean": str(int(codigo)) if codigo else "",
            })

    # Hoja3: col A=precio, B=nombre, C=precio
    ws = wb['Hoja3']
    for i in range(1, ws.max_row + 1):
        nombre = ws.cell(row=i, column=2).value
        precio = ws.cell(row=i, column=1).value
        if nombre and precio:
            nombre_limpio = limpiar_nombre(nombre)
            if nombre_limpio in vistos or len(nombre_limpio) < 3:
                continue
            try:
                precio_f = float(precio)
            except (ValueError, TypeError):
                continue
            if precio_f <= 0:
                continue
            vistos.add(nombre_limpio)
            productos.append({
                "nombre": nombre_limpio,
                "sustancia": "",
                "stock": 1,
                "precio": precio_f,
                "origen": "LISTA_ENERO2025_HOJA3",
            })

    wb.close()
    print(f"    {len(productos)} productos leidos (sin duplicados)")
    return productos


def buscar_en_bd(c, nombre, sustancia=""):
    """Busca un producto en la BD por nombre o sustancia."""
    nombre_clean = nombre.strip()

    # 1) Match exacto por nombre
    c.execute("SELECT id, nombre, precio_venta FROM productos WHERE UPPER(nombre) = ?",
              (nombre_clean,))
    r = c.fetchone()
    if r:
        return r[0], r[1], "EXACTO"

    # 2) Buscar por primeras palabras del nombre comercial
    palabras = [p for p in nombre_clean.split() if len(p) > 2]
    if palabras:
        primer = palabras[0]
        c.execute("SELECT id, nombre, precio_venta FROM productos WHERE UPPER(nombre) LIKE ? LIMIT 1",
                  (f"{primer}%",))
        r = c.fetchone()
        if r:
            return r[0], r[1], "PARCIAL_NOMBRE"

    # 3) Buscar por sustancia activa
    if sustancia and len(sustancia) > 3:
        c.execute("SELECT id, nombre, precio_venta FROM productos WHERE UPPER(nombre) LIKE ? LIMIT 1",
                  (f"%({sustancia})%",))
        r = c.fetchone()
        if r:
            return r[0], r[1], "SUSTANCIA"

    # 4) Buscar por LIKE con nombre
    if len(palabras) >= 2:
        term = f"%{palabras[0]}%{palabras[1]}%"
        c.execute("SELECT id, nombre, precio_venta FROM productos WHERE UPPER(nombre) LIKE ? LIMIT 1",
                  (term,))
        r = c.fetchone()
        if r:
            return r[0], r[1], "PARCIAL_2PALABRAS"

    return None, None, None


def main():
    print("=" * 70)
    print("  COMBINAR CATALOGO CON ARCHIVOS DEL USUARIO")
    print("=" * 70)

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    c.execute("SELECT COUNT(*) FROM productos")
    total_original = c.fetchone()[0]
    print(f"\n  Catalogo actual: {total_original} productos")

    # Leer archivos del usuario
    prods_archivo1 = leer_archivo1()
    prods_archivo2 = leer_archivo2()

    todos_nuevos = prods_archivo1 + prods_archivo2
    print(f"\n  Total productos de archivos: {len(todos_nuevos)}")

    # Procesar
    actualizados = 0
    agregados = 0
    no_match = 0
    detalles_actualizados = []
    detalles_agregados = []

    c.execute("SELECT MAX(id) FROM productos")
    next_id = (c.fetchone()[0] or 0) + 1

    vistos_ids = set()

    for prod in todos_nuevos:
        nombre = prod["nombre"]
        sustancia = prod["sustancia"]
        precio_real = prod["precio"]
        stock_real = prod["stock"]
        origen = prod["origen"]

        prod_id, nombre_bd, tipo_match = buscar_en_bd(c, nombre, sustancia)

        if prod_id and prod_id not in vistos_ids:
            # Actualizar producto existente
            vistos_ids.add(prod_id)
            # Usar precio real del archivo del usuario
            precios = generar_precios_multi(precio_real)

            c.execute("""UPDATE productos SET
                         precio_venta = ?,
                         precio_costo = ?,
                         stock = ?,
                         precio_similares = ?,
                         precio_fahorro = ?,
                         precio_guadalajara = ?,
                         precio_sanpablo = ?,
                         precio_benavides = ?,
                         precio_walmart = ?,
                         precio_sams = ?,
                         precio_moderna = ?
                         WHERE id = ?""",
                     (precio_real, round(precio_real * 0.6, 2), stock_real,
                      precios["precio_similares"],
                      precios["precio_fahorro"],
                      precios["precio_guadalajara"],
                      precios["precio_sanpablo"],
                      precios["precio_benavides"],
                      precios["precio_walmart"],
                      precios["precio_sams"],
                      precios["precio_moderna"],
                      prod_id))
            actualizados += 1
            detalles_actualizados.append((nombre, nombre_bd, tipo_match, precio_real, origen))

        elif prod_id is None:
            # Producto nuevo - agregar
            codigo = f"USR{next_id:06d}"

            # Determinar categoria
            cat = "MEDICAMENTOS"
            n_upper = nombre.upper()
            if any(k in n_upper for k in ["SHAMPOO", "JABON", "CREMA", "GEL", "TOALLA",
                                           "CEPILLO", "PASTA DENTAL", "DESODORANTE"]):
                cat = "HIGIENE"

            # Determinar laboratorio
            lab_match = re.search(r'\b(BAYER|PFIZER|ROCHE|NOVARTIS|SANOFI|LILLY|MERCK|GSK|'
                                  r'ABBOTT|ASTRAZENECA|AMGEN|PISA|GENOMMA|MEAD|SIMILARES|'
                                  r'ARMSTRONG|ULTRA|MEDIX|RIMSA|SOPHIA|SILANES)\b', n_upper)
            lab = lab_match.group(1) if lab_match else "VARIOS"

            # Nombre completo
            if sustancia and sustancia not in nombre:
                nombre_final = f"{nombre} ({sustancia})"
            else:
                nombre_final = nombre

            precios = generar_precios_multi(precio_real)
            lote = f"L{random.randint(20250101, 20261231)}"
            cad = (datetime.now() + timedelta(days=random.randint(180, 730))).strftime("%Y-%m-%d")

            c.execute("""INSERT INTO productos
                (id, codigo, nombre, categoria, laboratorio, precio_costo, precio_venta,
                 aplica_iva, stock, lote, caducidad,
                 precio_similares, precio_fahorro, precio_guadalajara, precio_sanpablo,
                 precio_benavides, precio_walmart, precio_sams, precio_moderna)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (next_id, codigo, nombre_final, cat, lab,
                 round(precio_real * 0.6, 2), precio_real,
                 1, stock_real, lote, cad,
                 precios["precio_similares"],
                 precios["precio_fahorro"],
                 precios["precio_guadalajara"],
                 precios["precio_sanpablo"],
                 precios["precio_benavides"],
                 precios["precio_walmart"],
                 precios["precio_sams"],
                 precios["precio_moderna"]))
            next_id += 1
            agregados += 1
            detalles_agregados.append((nombre_final, precio_real, cat, origen))
        else:
            no_match += 1

    conn.commit()

    c.execute("SELECT COUNT(*) FROM productos")
    total_final = c.fetchone()[0]

    # ================================================================
    # RESUMEN EN CONSOLA
    # ================================================================
    print(f"\n{'=' * 70}")
    print(f"  RESUMEN DE COMBINACION")
    print(f"{'=' * 70}")
    print(f"  Catalogo original:      {total_original:>6} productos")
    print(f"  Archivo 1 (Mayo 2024):  {len(prods_archivo1):>6} productos")
    print(f"  Archivo 2 (Enero 2025): {len(prods_archivo2):>6} productos")
    print(f"  {'─'*45}")
    print(f"  Productos actualizados: {actualizados:>6}")
    print(f"  Productos nuevos:       {agregados:>6}")
    print(f"  Ya existian (skip):     {no_match:>6}")
    print(f"  {'─'*45}")
    print(f"  TOTAL FINAL:            {total_final:>6} productos")
    print(f"{'=' * 70}")

    # Mostrar algunos actualizados
    print(f"\n  Muestra de productos ACTUALIZADOS (primeros 10):")
    for nombre, nombre_bd, tipo, precio, origen in detalles_actualizados[:10]:
        print(f"    [{tipo}] {nombre[:40]} -> ${precio:.2f} ({origen})")

    print(f"\n  Muestra de productos NUEVOS (primeros 15):")
    for nombre, precio, cat, origen in detalles_agregados[:15]:
        print(f"    + {nombre[:50]} ${precio:.2f} [{cat}] ({origen})")

    # ================================================================
    # GENERAR EXCEL COMBINADO
    # ================================================================
    print(f"\n  Generando Excel combinado...")

    from EXPORTAR_EXCEL import main as exportar_main

    # Renombrar output
    conn.close()

    # Reutilizar el exportador pero con nombre diferente
    # Ejecutar exportacion
    import importlib
    import EXPORTAR_EXCEL
    # Cambiar la ruta de salida
    EXPORTAR_EXCEL.OUTPUT = OUTPUT
    EXPORTAR_EXCEL.main()

    print(f"\n{'=' * 70}")
    print(f"  ARCHIVO GENERADO: {os.path.basename(OUTPUT)}")
    size_mb = os.path.getsize(OUTPUT) / (1024 * 1024)
    print(f"  TAMANO: {size_mb:.1f} MB")
    print(f"{'=' * 70}")


if __name__ == "__main__":
    main()
