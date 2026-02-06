"""
EXPORTAR CATALOGO COMPLETO A EXCEL
Genera archivo profesional con 4 hojas, formato, filtros y colores.
"""
import sqlite3
import os
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from datetime import datetime

DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "farmacia.db")
OUTPUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                      "CATALOGO_FARMACIAS_MADRID_COMPLETO.xlsx")

# Colores
AZUL = "4A90E2"
AZUL_OSCURO = "2C3E50"
VERDE = "27AE60"
NARANJA = "E67E22"
ROJO = "E74C3C"
GRIS_CLARO = "F2F2F2"
BLANCO = "FFFFFF"
AMARILLO_CLARO = "FFF9E6"

HEADER_FILL = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color=BLANCO, size=11)
TITLE_FONT = Font(name="Calibri", bold=True, color=AZUL_OSCURO, size=16)
SUBTITLE_FONT = Font(name="Calibri", bold=True, color=AZUL_OSCURO, size=12)
NORMAL_FONT = Font(name="Calibri", size=10)
MONEY_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", bold=True, size=10)
VERDE_FILL = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
ROJO_FILL = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
GRIS_FILL = PatternFill(start_color=GRIS_CLARO, end_color=GRIS_CLARO, fill_type="solid")
AMARILLO_FILL = PatternFill(start_color=AMARILLO_CLARO, end_color=AMARILLO_CLARO, fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

MONEY_FORMAT = '#,##0.00'
PCT_FORMAT = '0.0%'


def extraer_principio_activo(nombre):
    m = re.search(r'\(([^)]+)\)', nombre)
    return m.group(1) if m else ""


def extraer_presentacion(nombre):
    m = re.search(r'\)\s*(.+?)\s*-\s*\w', nombre)
    return m.group(1).strip() if m else ""


def apply_header(ws, row, cols, fill=None, font=None):
    f = fill or HEADER_FILL
    fn = font or HEADER_FONT
    for col_idx, text in enumerate(cols, 1):
        cell = ws.cell(row=row, column=col_idx, value=text)
        cell.fill = f
        cell.font = fn
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def auto_width(ws, min_w=8, max_w=45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_w), max_w)


def main():
    print("=" * 60)
    print("  EXPORTAR CATALOGO A EXCEL")
    print("=" * 60)

    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()

    wb = Workbook()

    # ================================================================
    # HOJA 1: CATALOGO COMPLETO
    # ================================================================
    print("  Generando hoja 'Catalogo Completo'...")
    ws1 = wb.active
    ws1.title = "Catalogo Completo"
    ws1.sheet_properties.tabColor = AZUL

    headers = [
        "Codigo", "Nombre Comercial", "Principio Activo", "Presentacion",
        "Laboratorio", "Categoria", "Stock",
        "Precio Costo", "Precio Venta",
        "Similares", "Del Ahorro", "Guadalajara", "San Pablo",
        "Benavides", "Walmart", "Sam's Club", "Moderna/Ahorrerra",
        "Precio Promedio", "Margen %"
    ]

    # Titulo
    ws1.merge_cells("A1:S1")
    title_cell = ws1.cell(row=1, column=1, value="CATALOGO FARMACIAS MADRID - COMPLETO")
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center")

    ws1.merge_cells("A2:S2")
    sub_cell = ws1.cell(row=2, column=1,
                        value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')} | "
                              f"Precios base: Farmacias Similares (API VTEX)")
    sub_cell.font = Font(name="Calibri", size=9, color="7F8C8D")
    sub_cell.alignment = Alignment(horizontal="center")

    apply_header(ws1, 4, headers)
    ws1.row_dimensions[4].height = 30

    c.execute("""SELECT codigo, nombre, laboratorio, categoria, stock,
                 precio_costo, precio_venta,
                 precio_similares, precio_fahorro, precio_guadalajara, precio_sanpablo,
                 precio_benavides, precio_walmart, precio_sams, precio_moderna
                 FROM productos ORDER BY categoria, nombre""")

    row_num = 5
    for row in c.fetchall():
        codigo, nombre, lab, cat, stock, pc, pv = row[:7]
        p_sim, p_fa, p_gd, p_sp, p_bn, p_wm, p_sm, p_md = row[7:]

        principio = extraer_principio_activo(nombre)
        presentacion = extraer_presentacion(nombre)

        precios = [p_sim or 0, p_fa or 0, p_gd or 0, p_sp or 0,
                   p_bn or 0, p_wm or 0, p_sm or 0, p_md or 0]
        precios_validos = [p for p in precios if p > 0]
        promedio = sum(precios_validos) / len(precios_validos) if precios_validos else 0

        pv_f = float(pv) if pv else 0
        pc_f = float(pc) if pc else 0
        margen = (pv_f - pc_f) / pv_f if pv_f > 0 else 0

        data = [
            codigo, nombre, principio, presentacion,
            lab, cat, stock or 0,
            pc_f, pv_f,
            float(p_sim or 0), float(p_fa or 0), float(p_gd or 0), float(p_sp or 0),
            float(p_bn or 0), float(p_wm or 0), float(p_sm or 0), float(p_md or 0),
            round(promedio, 2), margen
        ]

        for col_idx, val in enumerate(data, 1):
            cell = ws1.cell(row=row_num, column=col_idx, value=val)
            cell.font = NORMAL_FONT
            cell.border = THIN_BORDER

            # Formato moneda para columnas de precios (8-18)
            if 8 <= col_idx <= 18:
                cell.number_format = MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right")
            elif col_idx == 19:  # Margen
                cell.number_format = PCT_FORMAT
                cell.alignment = Alignment(horizontal="right")
                if margen >= 0.4:
                    cell.fill = VERDE_FILL
                elif margen < 0.2:
                    cell.fill = ROJO_FILL
            elif col_idx == 7:  # Stock
                cell.alignment = Alignment(horizontal="center")
                if (stock or 0) < 10:
                    cell.fill = ROJO_FILL
                    cell.font = Font(name="Calibri", size=10, color=ROJO)

            # Rayas alternas
            if row_num % 2 == 0 and col_idx not in (7, 19):
                cell.fill = GRIS_FILL

        row_num += 1

    # Filtros
    ws1.auto_filter.ref = f"A4:S{row_num - 1}"

    # Freeze panes
    ws1.freeze_panes = "A5"

    # Anchos personalizados
    widths = {1: 12, 2: 45, 3: 22, 4: 30, 5: 18, 6: 18, 7: 8,
              8: 12, 9: 12, 10: 12, 11: 12, 12: 13, 13: 12,
              14: 12, 15: 12, 16: 12, 17: 15, 18: 13, 19: 10}
    for col, w in widths.items():
        ws1.column_dimensions[get_column_letter(col)].width = w

    total_productos = row_num - 5
    print(f"    {total_productos} productos exportados")

    # ================================================================
    # HOJA 2: RESUMEN
    # ================================================================
    print("  Generando hoja 'Resumen'...")
    ws2 = wb.create_sheet("Resumen")
    ws2.sheet_properties.tabColor = VERDE

    # Titulo
    ws2.merge_cells("A1:F1")
    ws2.cell(row=1, column=1, value="RESUMEN DEL CATALOGO").font = TITLE_FONT
    ws2.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    ws2.merge_cells("A2:F2")
    ws2.cell(row=2, column=1,
             value=f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = \
        Font(name="Calibri", size=9, color="7F8C8D")
    ws2.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    # --- Estadisticas generales ---
    r = 4
    ws2.merge_cells(f"A{r}:F{r}")
    ws2.cell(row=r, column=1, value="ESTADISTICAS GENERALES").font = SUBTITLE_FONT
    r += 1

    c.execute("SELECT COUNT(*) FROM productos")
    total = c.fetchone()[0]
    c.execute("SELECT COUNT(DISTINCT categoria) FROM productos")
    num_cats = c.fetchone()[0]
    c.execute("SELECT COUNT(DISTINCT laboratorio) FROM productos")
    num_labs = c.fetchone()[0]
    c.execute("SELECT SUM(stock) FROM productos")
    total_stock = c.fetchone()[0] or 0

    stats = [
        ("Total de productos", total),
        ("Categorias", num_cats),
        ("Laboratorios", num_labs),
        ("Stock total (unidades)", total_stock),
    ]
    for label, val in stats:
        ws2.cell(row=r, column=1, value=label).font = BOLD_FONT
        cell = ws2.cell(row=r, column=2, value=val)
        cell.font = BOLD_FONT
        cell.number_format = '#,##0'
        cell.alignment = Alignment(horizontal="right")
        r += 1

    # --- Productos por categoria ---
    r += 1
    ws2.merge_cells(f"A{r}:F{r}")
    ws2.cell(row=r, column=1, value="PRODUCTOS POR CATEGORIA").font = SUBTITLE_FONT
    r += 1

    cat_headers = ["Categoria", "Productos", "Promedio PV", "Min PV", "Max PV", "Stock Total"]
    cat_fill = PatternFill(start_color=VERDE, end_color=VERDE, fill_type="solid")
    apply_header(ws2, r, cat_headers, fill=cat_fill)
    r += 1

    c.execute("""SELECT categoria, COUNT(*), ROUND(AVG(precio_venta),2),
                 ROUND(MIN(precio_venta),2), ROUND(MAX(precio_venta),2),
                 SUM(stock)
                 FROM productos GROUP BY categoria ORDER BY COUNT(*) DESC""")
    for row in c.fetchall():
        ws2.cell(row=r, column=1, value=row[0]).font = NORMAL_FONT
        ws2.cell(row=r, column=2, value=row[1]).font = NORMAL_FONT
        ws2.cell(row=r, column=2).number_format = '#,##0'
        for ci in range(3, 6):
            cell = ws2.cell(row=r, column=ci, value=float(row[ci - 1]))
            cell.number_format = MONEY_FORMAT
            cell.font = NORMAL_FONT
        ws2.cell(row=r, column=6, value=row[5] or 0).font = NORMAL_FONT
        ws2.cell(row=r, column=6).number_format = '#,##0'
        for ci in range(1, 7):
            ws2.cell(row=r, column=ci).border = THIN_BORDER
            if r % 2 == 0:
                ws2.cell(row=r, column=ci).fill = GRIS_FILL
        r += 1

    # --- Precios promedio por farmacia ---
    r += 1
    ws2.merge_cells(f"A{r}:F{r}")
    ws2.cell(row=r, column=1, value="PRECIO PROMEDIO POR FARMACIA").font = SUBTITLE_FONT
    r += 1

    farm_headers = ["Farmacia", "Precio Promedio", "Posicion"]
    farm_fill = PatternFill(start_color=NARANJA, end_color=NARANJA, fill_type="solid")
    apply_header(ws2, r, farm_headers, fill=farm_fill)
    r += 1

    farmacias_cols = [
        ("Farmacias Similares", "precio_similares"),
        ("Farmacias del Ahorro", "precio_fahorro"),
        ("Farmacias Guadalajara", "precio_guadalajara"),
        ("Farmacias San Pablo", "precio_sanpablo"),
        ("Farmacias Benavides", "precio_benavides"),
        ("Walmart Pharmacy", "precio_walmart"),
        ("Sam's Club Pharmacy", "precio_sams"),
        ("La Moderna/Ahorrerra", "precio_moderna"),
    ]

    promedios_farm = []
    for nombre_farm, col in farmacias_cols:
        c.execute(f"SELECT ROUND(AVG({col}), 2) FROM productos WHERE {col} > 0")
        avg = c.fetchone()[0] or 0
        promedios_farm.append((nombre_farm, avg))

    promedios_farm.sort(key=lambda x: x[1])
    for pos, (nombre_farm, avg) in enumerate(promedios_farm, 1):
        ws2.cell(row=r, column=1, value=nombre_farm).font = NORMAL_FONT
        cell = ws2.cell(row=r, column=2, value=avg)
        cell.number_format = MONEY_FORMAT
        cell.font = BOLD_FONT
        pos_cell = ws2.cell(row=r, column=3, value=f"#{pos}")
        pos_cell.font = BOLD_FONT
        pos_cell.alignment = Alignment(horizontal="center")
        if pos == 1:
            pos_cell.fill = VERDE_FILL
        elif pos == len(promedios_farm):
            pos_cell.fill = ROJO_FILL
        for ci in range(1, 4):
            ws2.cell(row=r, column=ci).border = THIN_BORDER
        r += 1

    # --- Top laboratorios ---
    r += 1
    ws2.merge_cells(f"A{r}:F{r}")
    ws2.cell(row=r, column=1, value="TOP 15 LABORATORIOS").font = SUBTITLE_FONT
    r += 1

    lab_headers = ["Laboratorio", "Productos", "Promedio PV"]
    lab_fill = PatternFill(start_color=AZUL_OSCURO, end_color=AZUL_OSCURO, fill_type="solid")
    apply_header(ws2, r, lab_headers, fill=lab_fill)
    r += 1

    c.execute("""SELECT laboratorio, COUNT(*), ROUND(AVG(precio_venta), 2)
                 FROM productos GROUP BY laboratorio
                 ORDER BY COUNT(*) DESC LIMIT 15""")
    for row in c.fetchall():
        ws2.cell(row=r, column=1, value=row[0]).font = NORMAL_FONT
        ws2.cell(row=r, column=2, value=row[1]).font = NORMAL_FONT
        ws2.cell(row=r, column=2).number_format = '#,##0'
        ws2.cell(row=r, column=3, value=float(row[2])).font = NORMAL_FONT
        ws2.cell(row=r, column=3).number_format = MONEY_FORMAT
        for ci in range(1, 4):
            ws2.cell(row=r, column=ci).border = THIN_BORDER
        r += 1

    # Anchos
    for col, w in {1: 28, 2: 16, 3: 14, 4: 12, 5: 12, 6: 14}.items():
        ws2.column_dimensions[get_column_letter(col)].width = w

    ws2.freeze_panes = "A4"

    # ================================================================
    # HOJA 3: TOP 100 MAS VENDIDOS
    # ================================================================
    print("  Generando hoja 'Top 100 Mas Vendidos'...")
    ws3 = wb.create_sheet("Top 100 Mas Vendidos")
    ws3.sheet_properties.tabColor = NARANJA

    ws3.merge_cells("A1:H1")
    ws3.cell(row=1, column=1, value="TOP 100 PRODUCTOS MAS VENDIDOS").font = TITLE_FONT
    ws3.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    # Verificar si hay datos de ventas
    c.execute("SELECT COUNT(*) FROM detalle_ventas")
    hay_ventas = c.fetchone()[0] > 0

    if hay_ventas:
        top_headers = ["#", "Codigo", "Producto", "Unidades Vendidas",
                       "Total Vendido", "Precio Unitario", "Stock Actual", "Categoria"]
        top_fill = PatternFill(start_color=NARANJA, end_color=NARANJA, fill_type="solid")
        apply_header(ws3, 3, top_headers, fill=top_fill)

        c.execute("""SELECT p.codigo, p.nombre, SUM(dv.cantidad) as unidades,
                     SUM(dv.subtotal) as total, p.precio_venta, p.stock, p.categoria
                     FROM detalle_ventas dv
                     JOIN productos p ON dv.producto_id = p.id
                     GROUP BY dv.producto_id
                     ORDER BY unidades DESC
                     LIMIT 100""")

        r = 4
        for pos, row in enumerate(c.fetchall(), 1):
            ws3.cell(row=r, column=1, value=pos).font = BOLD_FONT
            ws3.cell(row=r, column=1).alignment = Alignment(horizontal="center")
            ws3.cell(row=r, column=2, value=row[0]).font = NORMAL_FONT
            ws3.cell(row=r, column=3, value=row[1][:55]).font = NORMAL_FONT
            ws3.cell(row=r, column=4, value=row[2]).font = BOLD_FONT
            ws3.cell(row=r, column=4).number_format = '#,##0'
            ws3.cell(row=r, column=5, value=float(row[3])).font = NORMAL_FONT
            ws3.cell(row=r, column=5).number_format = MONEY_FORMAT
            ws3.cell(row=r, column=6, value=float(row[4])).font = NORMAL_FONT
            ws3.cell(row=r, column=6).number_format = MONEY_FORMAT
            ws3.cell(row=r, column=7, value=row[5] or 0).font = NORMAL_FONT
            ws3.cell(row=r, column=8, value=row[6]).font = NORMAL_FONT
            for ci in range(1, 9):
                ws3.cell(row=r, column=ci).border = THIN_BORDER
                if r % 2 == 0:
                    ws3.cell(row=r, column=ci).fill = GRIS_FILL
            # Top 3 highlight
            if pos <= 3:
                for ci in range(1, 9):
                    ws3.cell(row=r, column=ci).fill = AMARILLO_FILL
            r += 1

        ws3.auto_filter.ref = f"A3:H{r - 1}"
    else:
        ws3.merge_cells("A3:H3")
        ws3.cell(row=3, column=1,
                 value="No hay datos de ventas registrados. "
                       "Realiza ventas en el sistema para ver este reporte.").font = \
            Font(name="Calibri", size=11, color=ROJO, italic=True)

        # Mostrar top por stock como alternativa
        ws3.merge_cells("A5:H5")
        ws3.cell(row=5, column=1,
                 value="ALTERNATIVA: TOP 100 PRODUCTOS CON MAYOR STOCK").font = SUBTITLE_FONT

        alt_headers = ["#", "Codigo", "Producto", "Stock",
                       "Precio Venta", "Valor Inventario", "Categoria", "Laboratorio"]
        alt_fill = PatternFill(start_color="5D6D7E", end_color="5D6D7E", fill_type="solid")
        apply_header(ws3, 7, alt_headers, fill=alt_fill)

        c.execute("""SELECT codigo, nombre, stock, precio_venta, categoria, laboratorio
                     FROM productos ORDER BY stock DESC LIMIT 100""")

        r = 8
        for pos, row in enumerate(c.fetchall(), 1):
            stock_val = row[2] or 0
            pv_val = float(row[3]) if row[3] else 0
            ws3.cell(row=r, column=1, value=pos).font = BOLD_FONT
            ws3.cell(row=r, column=1).alignment = Alignment(horizontal="center")
            ws3.cell(row=r, column=2, value=row[0]).font = NORMAL_FONT
            ws3.cell(row=r, column=3, value=row[1][:55]).font = NORMAL_FONT
            ws3.cell(row=r, column=4, value=stock_val).font = BOLD_FONT
            ws3.cell(row=r, column=4).number_format = '#,##0'
            ws3.cell(row=r, column=5, value=pv_val).font = NORMAL_FONT
            ws3.cell(row=r, column=5).number_format = MONEY_FORMAT
            ws3.cell(row=r, column=6, value=round(stock_val * pv_val, 2)).font = NORMAL_FONT
            ws3.cell(row=r, column=6).number_format = MONEY_FORMAT
            ws3.cell(row=r, column=7, value=row[4]).font = NORMAL_FONT
            ws3.cell(row=r, column=8, value=row[5]).font = NORMAL_FONT
            for ci in range(1, 9):
                ws3.cell(row=r, column=ci).border = THIN_BORDER
                if r % 2 == 0:
                    ws3.cell(row=r, column=ci).fill = GRIS_FILL
            r += 1

        ws3.auto_filter.ref = f"A7:H{r - 1}"

    for col, w in {1: 5, 2: 12, 3: 50, 4: 16, 5: 14, 6: 16, 7: 14, 8: 18}.items():
        ws3.column_dimensions[get_column_letter(col)].width = w

    ws3.freeze_panes = "A4" if hay_ventas else "A8"

    # ================================================================
    # HOJA 4: COMPARATIVA DE PRECIOS
    # ================================================================
    print("  Generando hoja 'Comparativa Precios'...")
    ws4 = wb.create_sheet("Comparativa Precios")
    ws4.sheet_properties.tabColor = "9B59B6"

    ws4.merge_cells("A1:J1")
    ws4.cell(row=1, column=1, value="COMPARATIVA DE PRECIOS ENTRE FARMACIAS").font = TITLE_FONT
    ws4.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    ws4.merge_cells("A2:J2")
    ws4.cell(row=2, column=1,
             value="Analisis de diferencia de precios - Farmacia mas barata vs mas cara por producto").font = \
        Font(name="Calibri", size=9, color="7F8C8D")
    ws4.cell(row=2, column=1).alignment = Alignment(horizontal="center")

    comp_headers = ["Producto", "Categoria", "Mas Barata", "Precio Min",
                    "Mas Cara", "Precio Max", "Diferencia $", "Diferencia %",
                    "Precio Promedio", "Tu Precio"]
    comp_fill = PatternFill(start_color="9B59B6", end_color="9B59B6", fill_type="solid")
    apply_header(ws4, 4, comp_headers, fill=comp_fill)

    farmacias_map = {
        "precio_similares": "Similares",
        "precio_fahorro": "Del Ahorro",
        "precio_guadalajara": "Guadalajara",
        "precio_sanpablo": "San Pablo",
        "precio_benavides": "Benavides",
        "precio_walmart": "Walmart",
        "precio_sams": "Sam's Club",
        "precio_moderna": "Moderna",
    }

    c.execute("""SELECT nombre, categoria, precio_venta,
                 precio_similares, precio_fahorro, precio_guadalajara, precio_sanpablo,
                 precio_benavides, precio_walmart, precio_sams, precio_moderna
                 FROM productos WHERE precio_similares > 0
                 ORDER BY precio_venta DESC LIMIT 200""")

    r = 5
    for row in c.fetchall():
        nombre, cat, pv = row[0], row[1], float(row[2])
        precios = {}
        cols_list = list(farmacias_map.keys())
        for i, col_name in enumerate(cols_list):
            val = float(row[3 + i]) if row[3 + i] else 0
            if val > 0:
                precios[farmacias_map[col_name]] = val

        if not precios:
            continue

        min_farm = min(precios, key=precios.get)
        max_farm = max(precios, key=precios.get)
        min_p = precios[min_farm]
        max_p = precios[max_farm]
        diff = max_p - min_p
        diff_pct = diff / min_p if min_p > 0 else 0
        avg_p = sum(precios.values()) / len(precios)

        ws4.cell(row=r, column=1, value=nombre[:55]).font = NORMAL_FONT
        ws4.cell(row=r, column=2, value=cat).font = NORMAL_FONT
        ws4.cell(row=r, column=3, value=min_farm).font = Font(name="Calibri", size=10, color=VERDE)
        ws4.cell(row=r, column=4, value=min_p).number_format = MONEY_FORMAT
        ws4.cell(row=r, column=4).font = NORMAL_FONT
        ws4.cell(row=r, column=5, value=max_farm).font = Font(name="Calibri", size=10, color=ROJO)
        ws4.cell(row=r, column=6, value=max_p).number_format = MONEY_FORMAT
        ws4.cell(row=r, column=6).font = NORMAL_FONT
        ws4.cell(row=r, column=7, value=round(diff, 2)).number_format = MONEY_FORMAT
        ws4.cell(row=r, column=7).font = NORMAL_FONT
        ws4.cell(row=r, column=8, value=diff_pct).number_format = PCT_FORMAT
        ws4.cell(row=r, column=8).font = NORMAL_FONT
        if diff_pct > 0.3:
            ws4.cell(row=r, column=8).fill = ROJO_FILL
        ws4.cell(row=r, column=9, value=round(avg_p, 2)).number_format = MONEY_FORMAT
        ws4.cell(row=r, column=9).font = NORMAL_FONT
        ws4.cell(row=r, column=10, value=pv).number_format = MONEY_FORMAT
        ws4.cell(row=r, column=10).font = BOLD_FONT

        for ci in range(1, 11):
            ws4.cell(row=r, column=ci).border = THIN_BORDER
            if r % 2 == 0:
                ws4.cell(row=r, column=ci).fill = GRIS_FILL

        r += 1

    ws4.auto_filter.ref = f"A4:J{r - 1}"
    ws4.freeze_panes = "A5"

    for col, w in {1: 50, 2: 18, 3: 14, 4: 14, 5: 14, 6: 14, 7: 14, 8: 12, 9: 14, 10: 14}.items():
        ws4.column_dimensions[get_column_letter(col)].width = w

    # ================================================================
    # GUARDAR
    # ================================================================
    print(f"\n  Guardando {OUTPUT}...")
    wb.save(OUTPUT)

    size_mb = os.path.getsize(OUTPUT) / (1024 * 1024)
    print(f"\n{'=' * 60}")
    print(f"  ARCHIVO GENERADO: {os.path.basename(OUTPUT)}")
    print(f"  TAMANO: {size_mb:.1f} MB")
    print(f"  UBICACION: {OUTPUT}")
    print(f"{'=' * 60}")
    print(f"  Hojas:")
    print(f"    1. Catalogo Completo  ({total_productos} productos)")
    print(f"    2. Resumen            (estadisticas)")
    print(f"    3. Top 100 Mas Vendidos")
    print(f"    4. Comparativa Precios")
    print(f"{'=' * 60}")


if __name__ == "__main__":
    main()
