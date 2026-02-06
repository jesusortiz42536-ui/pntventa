import openpyxl

archivos = [
    (r'C:\Users\chule\Downloads\LISTA FARMACIA 06MAYO2024 (1).xlsx', 'ARCHIVO 1: LISTA FARMACIA 06MAYO2024'),
    (r'C:\Users\chule\Downloads\listaz enero 2025.xlsx', 'ARCHIVO 2: listaz enero 2025'),
]

for path, titulo in archivos:
    print('=' * 70)
    print(f'  {titulo}')
    print('=' * 70)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f'  Hojas: {wb.sheetnames}')
        for sn in wb.sheetnames:
            ws = wb[sn]
            print(f'\n  --- Hoja: {sn} ---')
            print(f'  Max fila: {ws.max_row}, Max col: {ws.max_column}')
            # Primeras 8 filas
            for i in range(1, min(9, ws.max_row + 1)):
                vals = []
                for j in range(1, min(16, ws.max_column + 1)):
                    v = ws.cell(row=i, column=j).value
                    vals.append(v)
                print(f'  Fila {i}: {vals}')
            # Contar filas con datos
            count = 0
            for i in range(2, ws.max_row + 1):
                has_data = False
                for j in range(1, ws.max_column + 1):
                    if ws.cell(row=i, column=j).value is not None:
                        has_data = True
                        break
                if has_data:
                    count += 1
            print(f'  Total filas con datos: {count}')
        wb.close()
    except Exception as e:
        print(f'  ERROR: {e}')
    print()
